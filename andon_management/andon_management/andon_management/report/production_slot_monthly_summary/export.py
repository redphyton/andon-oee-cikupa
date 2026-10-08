# File: andon_management/andon_management/report/production_slot_monthly_summary/export.py

import frappe
import json
import datetime
import io
import base64  # <-- TAMBAHKAN IMPORT INI
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

@frappe.whitelist()
def export_production_report(filters):
    if isinstance(filters, str):
        filters = json.loads(filters)
    filters = frappe._dict(filters)
    
    # 1. Dapatkan data laporan
    from .production_slot_monthly_summary import execute as get_report_data
    columns, data = get_report_data(filters)

    # 2. Buat Workbook langsung dengan Openpyxl
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Laporan Produksi'
    
    # Definisikan Styles
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    weekend_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    highlight_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    highlight_font = Font(color="9C0006", bold=True)
    grand_total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    right_align = Alignment(horizontal='right')

    # 3. Buat Header Kompleks
    year = int(filters.get('year'))
    month = int(filters.get('month'))
    
    header_row_1 = ['Mesin', 'Deskripsi Slot', 'Part No']
    header_row_2 = ['', '', '']
    
    date_cols = sorted(list(set([c['label'].split(' ')[0] for c in columns if c['fieldname'].startswith('plan_')])))
    
    for day in date_cols:
        header_row_1.extend([day, ''])
        header_row_2.extend(['Plan', 'Actual'])

    header_row_1.extend(['Total Plan', 'Total Actual'])
    header_row_2.extend(['', ''])

    worksheet.append(header_row_1)
    worksheet.append(header_row_2)

    # Atur merge untuk header
    worksheet.merge_cells('A1:A2')
    worksheet.merge_cells('B1:B2')
    worksheet.merge_cells('C1:C2')
    
    col_idx = 4
    for day in date_cols:
        start_col = get_column_letter(col_idx)
        end_col = get_column_letter(col_idx + 1)
        worksheet.merge_cells(f'{start_col}1:{end_col}1')
        col_idx += 2
        
    worksheet.merge_cells(f'{get_column_letter(col_idx)}1:{get_column_letter(col_idx)}2')
    worksheet.merge_cells(f'{get_column_letter(col_idx+1)}1:{get_column_letter(col_idx+1)}2')

    # 4. Tulis Data dan Terapkan Style
    for row_idx, row_data in enumerate(data, start=3):
        row_values = [row_data.get(c['fieldname'], '') for c in columns]
        worksheet.append(row_values)
        
        is_grand_total = 'GRAND TOTAL' in str(row_data.get('slot_no', ''))

        for col_idx, col_def in enumerate(columns, start=1):
            cell = worksheet.cell(row=row_idx, column=col_idx)

            if col_def['fieldname'].startswith(('plan_', 'actual_', 'total_')):
                cell.alignment = right_align

            if is_grand_total:
                cell.fill = grand_total_fill
                cell.font = Font(bold=True)
                continue

            if col_def['fieldname'].startswith(('plan_', 'actual_')):
                day_str = col_def['fieldname'].split('_')[1]
                
                try:
                    dt = datetime.date(year, month, int(day_str))
                    if dt.weekday() >= 5: # 5=Sabtu, 6=Minggu
                        cell.fill = weekend_fill
                except ValueError:
                    pass

                if col_def['fieldname'].startswith('actual_'):
                    plan_val = row_data.get(f"plan_{day_str}", 0)
                    actual_val = row_data.get(f"actual_{day_str}", 0)
                    if plan_val > 0 and actual_val < plan_val:
                        cell.fill = highlight_fill
                        cell.font = highlight_font
    
    for row in worksheet['1:2']:
        for cell in row:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            try:
                day = int(cell.value)
                dt = datetime.date(year, month, day)
                if dt.weekday() >= 5:
                    cell.fill = weekend_fill
            except (ValueError, TypeError):
                pass

    # 5. Simpan file ke memory stream
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    
    # --- PERUBAHAN DI SINI ---
    # 6. Kembalikan file sebagai base64 menggunakan metode yang benar
    encoded_content = base64.b64encode(stream.read()).decode('utf-8')
    
    return {
        "filename": f"Laporan Produksi Bulanan-{year}-{month}.xlsx",
        "filecontent": encoded_content
    }